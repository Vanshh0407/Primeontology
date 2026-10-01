"""PDF / DOCX / PPTX / TXT / MD / HTML -> sections [{page, heading, text}]."""
import io
import os
import re

from .common import IngestError
from .tabular import _decode

_HEADING = re.compile(
    r"^\s*(?:(?:ARTICLE|SECTION|CLAUSE|SCHEDULE|EXHIBIT)\s+[\dIVXA-Z]+[.:)-]?\s*.{0,80}"
    r"|\d+(?:\.\d+){0,3}[.)]?\s+[A-Z][^\n]{2,80}"
    r"|[A-Z][A-Z &/,'-]{4,70})\s*$"
)


def _is_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ";", ",")) and not re.match(r"^\d", line):
        return False
    if re.match(r"^\d+(\.\d+)*[.)]?\s+[A-Z]", line) and len(line.split()) > 12:
        return False
    return bool(_HEADING.match(line))


def split_sections(text: str, page: int | None = None) -> list[dict]:
    sections, heading, buf = [], None, []

    def flush():
        body = "\n".join(buf).strip()
        if body or heading:
            sections.append({"page": page, "heading": heading, "text": body})

    for raw in text.splitlines():
        md = re.match(r"^#{1,6}\s+(.*)$", raw)
        if md or _is_heading(raw):
            flush()
            heading, buf = (md.group(1) if md else raw).strip(), []
        else:
            buf.append(raw)
    flush()
    return sections


OCR_MAX_PAGES = 40  # bound the CPU time spent on one upload


def ocr_available() -> bool:
    if os.environ.get("PRIME_ONTOLOGY_OCR", "auto").lower() == "off":
        return False
    try:
        import pypdfium2  # noqa: F401
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except Exception:
        return False


_ocr_engine = None


def ocr_pages(data: bytes, page_numbers: list[int]) -> dict[int, str]:
    """Render PDF pages to images and run OCR (RapidOCR/ONNX, fully local). Returns {page_no: text} (1-based)."""
    global _ocr_engine
    import numpy as np
    import pypdfium2 as pdfium
    from rapidocr_onnxruntime import RapidOCR

    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    pdf = pdfium.PdfDocument(data)
    out = {}
    for n in page_numbers[:OCR_MAX_PAGES]:
        img = pdf[n - 1].render(scale=2.0).to_pil().convert("RGB")
        result, _ = _ocr_engine(np.array(img))
        # result rows: [box, text, score]; order top-to-bottom then left-to-right, group into lines by y
        rows = sorted(((r[0][0][1], r[0][0][0], r[1]) for r in (result or [])), key=lambda x: (round(x[0] / 14), x[1]))
        lines, cur, last = [], [], None
        for y, x, text in rows:
            key = round(y / 14)
            if last is not None and key != last:
                lines.append(" ".join(cur))
                cur = []
            cur.append(text)
            last = key
        if cur:
            lines.append(" ".join(cur))
        out[n] = "\n".join(lines)
    return out


def parse_pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    except Exception as e:
        raise IngestError(f"Cannot read PDF: {e}") from e
    warnings = []
    empty = [n for n, t in pages if len(t.strip()) < 10]
    if empty:
        if ocr_available():
            try:
                done = ocr_pages(data, empty)
                pages = [(n, done.get(n, t) if n in done and len(done[n].strip()) > len(t.strip()) else t) for n, t in pages]
                recovered = [n for n in done if done[n].strip()]
                warnings.append(f"OCR was applied to {len(recovered)} scanned page(s) {recovered[:10]}; recognised text may contain errors — review the evidence.")
                if len(empty) > OCR_MAX_PAGES:
                    warnings.append(f"Only the first {OCR_MAX_PAGES} scanned pages were OCR'd.")
            except Exception as e:  # OCR is best-effort; never lose the already-extracted text
                warnings.append(f"OCR failed ({type(e).__name__}); {len(empty)} page(s) have no text.")
        elif not any(t.strip() for _, t in pages):
            raise IngestError("PDF has no extractable text (likely a scanned image) and OCR is not available on this server. "
                              "Install the optional OCR packages (pip install pypdfium2 rapidocr-onnxruntime) or upload a searchable PDF.")
        else:
            warnings.append(f"{len(empty)} page(s) had no extractable text (possible scans, OCR unavailable): {empty[:10]}")
    if not any(t.strip() for _, t in pages):
        raise IngestError("PDF has no extractable text, even after OCR.")
    sections = []
    carry = None
    for n, text in pages:
        for s in split_sections(text, n):
            if carry and not s["heading"]:
                s["heading"] = carry
            sections.append(s)
            carry = s["heading"] or carry
    return {"sections": sections, "warnings": warnings, "pages": len(pages)}


def parse_docx(data: bytes) -> dict:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f"Cannot read DOCX: {e}") from e
    sections, heading, buf = [], None, []

    def flush():
        if buf or heading:
            sections.append({"page": None, "heading": heading, "text": "\n".join(buf).strip()})

    for p in d.paragraphs:
        if p.style is not None and p.style.name.lower().startswith(("heading", "title")) and p.text.strip():
            flush()
            heading, buf = p.text.strip(), []
        elif _is_heading(p.text) and p.runs and all(r.bold for r in p.runs if r.text.strip()):
            flush()
            heading, buf = p.text.strip(), []
        elif p.text.strip():
            buf.append(p.text.strip())
    for t in d.tables:
        for row in t.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells))
    flush()
    return {"sections": sections, "warnings": [], "pages": None}


def parse_pptx(data: bytes) -> dict:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f"Cannot read PPTX: {e}") from e
    sections = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else None
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        sections.append({"page": i, "heading": title or f"Slide {i}", "text": "\n".join(texts).strip()})
    return {"sections": sections, "warnings": [], "pages": len(sections)}


def parse_html(data: bytes) -> dict:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    sections, heading, buf = [], None, []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "pre"]):
        txt = el.get_text(" ", strip=True)
        if not txt:
            continue
        if el.name.startswith("h"):
            if buf or heading:
                sections.append({"page": None, "heading": heading, "text": "\n".join(buf)})
            heading, buf = txt, []
        else:
            buf.append(txt)
    if buf or heading:
        sections.append({"page": None, "heading": heading, "text": "\n".join(buf)})
    return {"sections": sections, "warnings": [], "pages": None}


def parse_text(data: bytes) -> dict:
    return {"sections": split_sections(_decode(data)), "warnings": [], "pages": None}


def parse_document(ext: str, data: bytes) -> dict:
    fn = {".pdf": lambda: parse_pdf(data), ".docx": lambda: parse_docx(data), ".pptx": lambda: parse_pptx(data),
          ".html": lambda: parse_html(data), ".htm": lambda: parse_html(data),
          ".txt": lambda: parse_text(data), ".md": lambda: parse_text(data), ".markdown": lambda: parse_text(data)}[ext]
    res = fn()
    if not any(s["text"].strip() for s in res["sections"]):
        raise IngestError("No text could be extracted from the document.")
    return res

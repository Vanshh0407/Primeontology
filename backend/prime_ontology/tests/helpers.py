import io
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
SAMPLES = ROOT / "demo" / "samples"
DEMO_SQL = ROOT / "demo" / "demo_schema.sql"


def sample(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


def make_pdf(pages: list[str]) -> bytes:
    """Minimal text PDF (one text object per page) for tests."""
    objs = []

    def add(b):
        objs.append(b)
        return len(objs)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    add(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, text in enumerate(pages):
        page_no, content_no = 3 + i * 2, 4 + i * 2
        add(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_no} 0 R "
            f"/Resources << /Font << /F1 {3 + len(pages) * 2} 0 R >> >> >>".encode())
        lines = text.split("\n")
        stream = "BT /F1 11 Tf 40 750 Td 14 TL " + " ".join(
            f"({l.replace('(', '[').replace(')', ']')}) Tj T*" for l in lines) + " ET"
        add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, b in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + b + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for o in offsets:
        out.write(f"{o:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    import docx

    d = docx.Document()
    for style, text in paragraphs:
        if style.startswith("Heading"):
            d.add_heading(text, level=int(style[-1]))
        else:
            d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_pptx(slides: list[tuple[str, str]]) -> bytes:
    from pptx import Presentation

    p = Presentation()
    for title, body in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


CONTRACT_TEXT = sample("services_agreement.txt").decode() if (SAMPLES / "services_agreement.txt").exists() else ""


def jpost(client, url, data=None, **kw):
    return client.post(url, json.dumps(data or {}), content_type="application/json", **kw)


def jput(client, url, data=None, **kw):
    return client.put(url, json.dumps(data or {}), content_type="application/json", **kw)
